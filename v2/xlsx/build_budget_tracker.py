# -*- coding: utf-8 -*-
"""Build: Personal Budget & Savings Tracker 2026 (XLSX)"""
import os, random
from datetime import date, timedelta
import calendar

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule, DataBarRule

# ----------------------------------------------------------------- STYLE
NAVY = "1F2A44"
TEAL = "0FA3B1"
LIGHT = "EAF4F4"
GOLD = "F4B942"
WHITE = "FFFFFF"
GREY = "8D99AE"
BAND = "F7FAFC"

FONT = "Calibri"
thin = Side(style="thin", color="D5DDE5")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

MONEY = "$#,##0.00"
DATEF = "dd-mmm-yyyy"
PCT = "0.0%"

TODAY = date(2026, 9, 27)
MONTHS_SHORT = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
                "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def hdr(ws, row, headers, start=1, fill=NAVY):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start + i, value=h)
        c.font = Font(name=FONT, bold=True, color=WHITE, size=11)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[row].height = 24


def band(ws, row, c1, c2, text, fill=NAVY, color=WHITE, size=11):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    c = ws.cell(row=row, column=c1, value=text)
    c.font = Font(name=FONT, bold=True, color=color, size=size)


def box(ws, r1, r2, c1, c2, money=(), dates=(), pct=(), ints=(), zebra=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if c in money:
                cell.number_format = MONEY
            elif c in dates:
                cell.number_format = DATEF
            elif c in pct:
                cell.number_format = PCT
            elif c in ints:
                cell.number_format = "#,##0"
            if zebra and (r - r1) % 2 == 1 and cell.fill.patternType is None:
                cell.fill = PatternFill("solid", fgColor=BAND)


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def sheet(wb, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    return ws


def start_here(ws, title, tagline, lines):
    widths(ws, {"A": 3, "B": 100})
    ws["B2"] = title
    ws["B2"].font = Font(name=FONT, bold=True, size=20, color=NAVY)
    ws["B3"] = tagline
    ws["B3"].font = Font(name=FONT, size=12, color=TEAL, italic=True)
    r = 5
    for line in lines:
        c = ws.cell(row=r, column=2, value=line)
        if line.isupper() and len(line) > 3:
            c.font = Font(name=FONT, bold=True, size=12, color=TEAL)
        else:
            c.font = Font(name=FONT, size=11, color="2B2D42")
        r += 1


def kpi(ws, row, col, label, formula, fmt=MONEY):
    lab = ws.cell(row=row, column=col, value=label)
    lab.font = Font(name=FONT, bold=True, size=10, color=WHITE)
    lab.fill = PatternFill("solid", fgColor=NAVY)
    lab.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lab.border = border
    val = ws.cell(row=row + 1, column=col, value=formula)
    val.font = Font(name=FONT, bold=True, size=17, color=TEAL)
    val.fill = PatternFill("solid", fgColor=LIGHT)
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.border = border
    val.number_format = fmt
    ws.row_dimensions[row].height = 28
    ws.row_dimensions[row + 1].height = 30


def dv_list(ws, options, rng, error="Pick a value from the list"):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = error
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def status_cf(ws, rng, rules):
    for op, fill, fc, bold in rules:
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="equal", formula=[f'"{op}"'],
                       fill=PatternFill("solid", fgColor=fill),
                       font=Font(name=FONT, color=fc, bold=bold)))


# ----------------------------------------------------------------- CATEGORIES
FIXED = [
    ("Rent", 1450, "Needs"),
    ("Utilities", 170, "Needs"),
    ("Insurance", 145, "Needs"),
    ("Phone", 45, "Needs"),
    ("Subscriptions", 35, "Wants"),
    ("Transport", 70, "Needs"),
]
VARIABLE = [
    ("Groceries", 250, "Needs"),
    ("Dining", 90, "Wants"),
    ("Health", 45, "Needs"),
    ("Clothing", 60, "Wants"),
    ("Entertainment", 65, "Wants"),
    ("Education", 30, "Wants"),
    ("Misc", 55, "Wants"),
    ("Household", 50, "Needs"),
]
CATEGORIES = [c for c, _p, _b in FIXED] + [c for c, _p, _b in VARIABLE]
BUCKETS = {c: b for c, _p, b in FIXED + VARIABLE}

# ----------------------------------------------------------------- WORKBOOK
wb = openpyxl.Workbook()

# ================================================================= START HERE
ins = wb.active
ins.title = "Start Here"
ins.sheet_view.showGridLines = False
start_here(
    ins,
    "PERSONAL BUDGET & SAVINGS TRACKER 2026",
    "Monthly budget  ·  Expense log  ·  Savings goals  ·  Debt snowball  ·  Net worth",
    [
        "",
        "HOW TO USE THIS WORKBOOK (5 minutes to set up)",
        "",
        "1. Open MONTHLY BUDGET and type your plan into the light-blue cells (Jan to Dec).",
        "2. Set 'Months with actuals' in cell D4 and your monthly income in cell D41.",
        "3. Log every purchase in EXPENSE LOG - pick Category, Method and Bucket from the drop-downs.",
        "4. The ACTUAL block on MONTHLY BUDGET fills itself in with SUMIFS - never type there.",
        "5. Track progress on SAVINGS GOALS (progress bar + auto status) and DEBT SNOWBALL.",
        "6. Update NET WORTH once a month - assets in, liabilities in, the bottom row does the maths.",
        "7. The DASHBOARD is the one page to read: change 'Reporting month' in cell C4.",
        "",
        "RULES THAT KEEP IT CLEAN",
        "",
        "- Never delete rows - add a correcting entry with the opposite amount instead.",
        "- Use the drop-down lists (Category, Method, Bucket, Essential?) so the totals keep working.",
        "- Enter amounts as plain numbers: type 42.50 not $42.50.",
        "- Dates use dd-mmm-yyyy, e.g. 03-May-2026.",
        "- One row = one transaction. Rent and bills count as expenses, not transfers.",
        "- Red on the budget grid means you have spent more than the plan allows - fix it fast.",
        "",
        "THE 50 / 30 / 20 RULE",
        "",
        "1. 50% of take-home pay for Needs: rent, utilities, food, transport, insurance.",
        "2. 30% for Wants: eating out, hobbies, streaming, shopping.",
        "3. 20% for Savings and extra debt payments - pay this first, on payday.",
        "4. The summary block at the bottom of MONTHLY BUDGET scores you against those targets.",
    ],
)

# ============================================================== EXPENSE LOG
el = sheet(wb, "EXPENSE LOG")
e_headers = ["Date", "Category", "Subcategory", "Merchant", "Method", "Amount",
             "Essential?", "Bucket", "Notes"]
e_widths = [13, 16, 18, 26, 15, 13, 13, 12, 30]
for i, (h, w) in enumerate(zip(e_headers, e_widths), start=1):
    el.cell(row=1, column=i, value=h)
    el.column_dimensions[get_column_letter(i)].width = w
hdr(el, 1, e_headers)

random.seed(42)
fixed_monthly = [
    ("Rent", "Apartment", "Harbor Property Mgmt", "Bank transfer", 1450, "Yes", "Needs", 1),
    ("Utilities", "Electric", "City Power & Light", "Debit card", 94, "Yes", "Needs", 5),
    ("Phone", "Mobile plan", "Mint Mobile", "Debit card", 45, "Yes", "Needs", 8),
    ("Transport", "Metro pass", "City Transit", "Debit card", 62, "Yes", "Needs", 12),
]
rotating = [
    ("Subscriptions", "Streaming", "Netflix", "Credit card", 17.99, "No", "Wants", 4),
    ("Subscriptions", "Music", "Spotify", "Credit card", 11.99, "No", "Wants", 9),
    ("Groceries", "Weekly shop", "Trader Joe's", "Credit card", 118, "Yes", "Needs", 6),
    ("Groceries", "Top-up shop", "Corner Market", "Debit card", 64, "Yes", "Needs", 17),
    ("Dining", "Coffee + lunch", "Bluebird Cafe", "Credit card", 47, "No", "Wants", 11),
    ("Dining", "Dinner out", "Osteria Vero", "Credit card", 92, "No", "Wants", 21),
    ("Health", "Pharmacy", "Walgreens", "Debit card", 26, "Yes", "Needs", 14),
    ("Entertainment", "Cinema", "Rialto Cinema", "Credit card", 34, "No", "Wants", 19),
    ("Education", "Online course", "Skillshare", "PayPal", 29, "No", "Wants", 23),
    ("Clothing", "Work shirts", "Uniqlo", "Credit card", 78, "No", "Wants", 16),
    ("Insurance", "Health cover", "Bright Health", "Bank transfer", 145, "Yes", "Needs", 3),
    ("Misc", "Household bits", "Amazon", "Credit card", 41, "No", "Wants", 25),
    ("Entertainment", "Concert ticket", "The Fillmore", "Credit card", 85, "No", "Wants", 27),
    ("Groceries", "Farmers market", "Saturday market", "Cash", 52, "Yes", "Needs", 20),
    ("Utilities", "Internet", "Sonic Fiber", "Bank transfer", 59, "Yes", "Needs", 7),
    ("Household", "Detergent + bulbs", "Target", "Debit card", 36, "Yes", "Needs", 24),
    ("Dining", "Team lunch", "Cafe Norte", "Credit card", 58, "No", "Wants", 26),
    ("Misc", "Gift for Sam", "Etsy", "PayPal", 45, "No", "Wants", 13),
]
notes_map = {"Rent": "Standing order", "Insurance": "Monthly premium",
             "Education": "Course instalment", "Entertainment": "Birthday treat"}

log_rows = []
for month in range(1, 6):
    last = calendar.monthrange(2026, month)[1]
    entries = list(fixed_monthly)
    for i in range(6):
        entries.append(rotating[(month * 6 + i) % len(rotating)])
    for cat, sub, mer, meth, amt, ess, buck, day in entries:
        if cat in ("Rent", "Phone"):
            amount = float(amt)
        else:
            amount = round(amt * random.uniform(0.88, 1.16), 2)
        d = date(2026, month, min(day, last))
        note = notes_map.get(cat, "")
        if meth == "Cash" and not note:
            note = "No card trail"
        log_rows.append([d, cat, sub, mer, meth, amount, ess, buck, note])

log_rows.sort(key=lambda x: (x[0], x[1]))
assert len(log_rows) == 50, len(log_rows)
for i, vals in enumerate(log_rows):
    r = 2 + i
    for c, v in enumerate(vals, start=1):
        el.cell(row=r, column=c, value=v)
last_el = 1 + len(log_rows)
box(el, 2, last_el, 1, 9, money=(6,), dates=(1,))
for r in range(2, last_el + 1):
    for c in (7, 8):
        el.cell(row=r, column=c).alignment = Alignment(horizontal="center")
el.freeze_panes = "A2"
el.auto_filter.ref = f"A1:I{last_el}"

dv_list(el, CATEGORIES, "B2:B1000")
dv_list(el, ["Debit card", "Credit card", "Cash", "Bank transfer", "PayPal"], "E2:E1000")
dv_list(el, ["Yes", "No"], "G2:G1000")
dv_list(el, ["Needs", "Wants", "Savings"], "H2:H1000")
status_cf(el, "G2:G1000", [("Yes", "D4EDDA", "1E7E34", False),
                           ("No", "E2E3E5", "6C757D", False)])
status_cf(el, "H2:H1000", [("Needs", "D1ECF1", "0C5460", False),
                           ("Wants", "FFF3CD", "997404", False),
                           ("Savings", "D4EDDA", "1E7E34", True)])
el.conditional_formatting.add(
    f"F2:F{last_el}",
    DataBarRule(start_type="num", start_value=0, end_type="max", color=TEAL))

# =========================================================== MONTHLY BUDGET
mb = sheet(wb, "MONTHLY BUDGET")
widths(mb, {"A": 3, "B": 24, "C": 11, "D": 11, "Q": 14, "R": 14, "S": 14, "T": 14, "U": 11})
for i in range(5, 17):  # E..P = Jan..Dec
    mb.column_dimensions[get_column_letter(i)].width = 12

mb["B2"] = "MONTHLY BUDGET 2026"
mb["B2"].font = Font(name=FONT, bold=True, size=18, color=NAVY)
mb["B3"] = "Plan cells are editable (light blue). The ACTUAL block reads the EXPENSE LOG with SUMIFS."
mb["B3"].font = Font(name=FONT, size=11, color=TEAL, italic=True)

lb = mb.cell(row=4, column=2, value="Months with actuals:")
lb.font = Font(name=FONT, bold=True, size=11, color=GREY)
lb.border = border
mv = mb.cell(row=4, column=4, value=5)
mv.font = Font(name=FONT, bold=True, size=12, color=TEAL)
mv.fill = PatternFill("solid", fgColor=LIGHT)
mv.border = border
mv.alignment = Alignment(horizontal="center")
note = mb.cell(row=4, column=5, value="(update this when you log a new month of expenses)")
note.font = Font(name=FONT, size=10, color=GREY, italic=True)

# ---- plan block
band(mb, 6, 2, 17, "PLANNED SPEND BY MONTH - EDIT THESE CELLS")
hdr(mb, 7, ["Category", "Group", "Bucket"] + MONTHS_SHORT + ["Plan total"], start=2, fill=TEAL)
plan_rows = [(c, "Fixed", b, p) for c, p, b in FIXED] + \
            [(c, "Variable", b, p) for c, p, b in VARIABLE]
for i, (cat, grp, buck, plan) in enumerate(plan_rows):
    r = 8 + i
    mb.cell(row=r, column=2, value=cat)
    mb.cell(row=r, column=3, value=grp)
    mb.cell(row=r, column=4, value=buck)
    for m in range(12):
        mb.cell(row=r, column=5 + m, value=plan)
    mb.cell(row=r, column=17, value=f"=SUM(E{r}:P{r})")
box(mb, 8, 21, 2, 17, money=tuple(range(5, 18)))
for r in range(8, 22):
    mb.cell(row=r, column=2).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
    for c in (3, 4):
        mb.cell(row=r, column=c).alignment = Alignment(horizontal="center")
        mb.cell(row=r, column=c).font = Font(name=FONT, size=10, color=GREY)
    for c in range(5, 17):
        mb.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LIGHT)
        mb.cell(row=r, column=c).alignment = Alignment(horizontal="right")
    mb.cell(row=r, column=17).font = Font(name=FONT, bold=True, size=11, color=NAVY)
    mb.cell(row=r, column=17).alignment = Alignment(horizontal="right")

# ---- actual block
band(mb, 23, 2, 21, "ACTUAL SPEND BY MONTH - FORMULAS ONLY, DO NOT TYPE HERE")
hdr(mb, 24, ["Category", "Group", "Bucket"] + MONTHS_SHORT +
    ["Actual total", "Plan (12 mo)", "Plan to date", "Variance", "% used"], start=2, fill=TEAL)
EL = "'EXPENSE LOG'"
for i, (cat, grp, buck, plan) in enumerate(plan_rows):
    r = 25 + i
    mb.cell(row=r, column=2, value=cat)
    mb.cell(row=r, column=3, value=grp)
    mb.cell(row=r, column=4, value=buck)
    for m in range(12):
        col = 5 + m
        letter = get_column_letter(col)
        mb.cell(row=r, column=col,
                value=f'=SUMIFS({EL}!$F:$F,{EL}!$B:$B,$B{r},{EL}!$A:$A,">="&DATE(2026,{m + 1},1),'
                      f'{EL}!$A:$A,"<"&DATE(2026,{m + 2},1))')
    mb.cell(row=r, column=17, value=f"=SUM(E{r}:P{r})")
    mb.cell(row=r, column=18, value=f"=VLOOKUP($B{r},$B$8:$Q$21,16,FALSE)")
    mb.cell(row=r, column=19, value=f"=R{r}*$D$4/12")
    mb.cell(row=r, column=20, value=f"=S{r}-Q{r}")
    mb.cell(row=r, column=21, value=f"=IF(S{r}=0,0,Q{r}/S{r})")
box(mb, 25, 38, 2, 21, money=tuple(range(5, 21)), pct=(21,), zebra=True)
for r in range(25, 39):
    mb.cell(row=r, column=2).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
    for c in (3, 4):
        mb.cell(row=r, column=c).alignment = Alignment(horizontal="center")
        mb.cell(row=r, column=c).font = Font(name=FONT, size=10, color=GREY)
    for c in range(5, 21):
        mb.cell(row=r, column=c).alignment = Alignment(horizontal="right")
    mb.cell(row=r, column=21).alignment = Alignment(horizontal="center")
    for c in (17, 18, 20):
        mb.cell(row=r, column=c).font = Font(name=FONT, bold=True, size=11, color=NAVY)
mb.conditional_formatting.add(
    "U25:U38",
    CellIsRule(operator="greaterThan", formula=["1"],
               fill=PatternFill("solid", fgColor="F8D7DA"),
               font=Font(name=FONT, bold=True, color="A02020")))
mb.conditional_formatting.add(
    "U25:U38",
    CellIsRule(operator="between", formula=["0.9", "1"],
               fill=PatternFill("solid", fgColor="FFF3CD"),
               font=Font(name=FONT, bold=True, color="997404")))
mb.conditional_formatting.add(
    "U25:U38",
    CellIsRule(operator="lessThan", formula=["0.9"],
               font=Font(name=FONT, color="1E7E34")))
mb.conditional_formatting.add(
    "T25:T38",
    CellIsRule(operator="lessThan", formula=["0"],
               font=Font(name=FONT, bold=True, color="A02020")))
mb.conditional_formatting.add(
    "Q25:Q38",
    DataBarRule(start_type="num", start_value=0, end_type="max", color=TEAL))

# ---- 50 / 30 / 20 block
band(mb, 40, 2, 8, "50 / 30 / 20 RULE - SCORE THE LAST 5 MONTHS")
lb2 = mb.cell(row=41, column=2, value="Monthly income (after tax)")
lb2.font = Font(name=FONT, bold=True, size=11, color=GREY)
lb2.border = border
mb.cell(row=41, column=3).border = border
iv = mb.cell(row=41, column=4, value=4200)
iv.font = Font(name=FONT, bold=True, size=12, color=TEAL)
iv.fill = PatternFill("solid", fgColor=LIGHT)
iv.border = border
iv.number_format = MONEY
iv.alignment = Alignment(horizontal="right")
for c in range(5, 9):
    mb.cell(row=41, column=c).border = border
hdr(mb, 42, ["Bucket", "Target %", "Target $", "Actual $", "% of income", "Gap $", "Status"],
    start=2, fill=TEAL)
rules_rows = [
    ("Needs", 0.50, "=SUMIF({0}!$H:$H,$B43,{0}!$F:$F)".format(EL), "over"),
    ("Wants", 0.30, "=SUMIF({0}!$H:$H,$B44,{0}!$F:$F)".format(EL), "over"),
    ("Savings", 0.20, "=($D$41*$D$4)-E43-E44", "under"),
]
for i, (buck, target, actual_f, direction) in enumerate(rules_rows):
    r = 43 + i
    mb.cell(row=r, column=2, value=buck)
    mb.cell(row=r, column=3, value=target)
    mb.cell(row=r, column=4, value=f"=$D$41*$D$4*C{r}")
    mb.cell(row=r, column=5, value=actual_f)
    mb.cell(row=r, column=6, value=f"=IF($D$41*$D$4=0,0,E{r}/($D$41*$D$4))")
    mb.cell(row=r, column=7, value=f"=D{r}-E{r}")
    if direction == "over":
        mb.cell(row=r, column=8, value=f'=IF(E{r}<=D{r},"On target","Over target")')
    else:
        mb.cell(row=r, column=8, value=f'=IF(E{r}>=D{r},"On target","Behind plan")')
mb.cell(row=46, column=2, value="TOTAL")
mb.cell(row=46, column=3, value="=SUM(C43:C45)")
mb.cell(row=46, column=4, value="=SUM(D43:D45)")
mb.cell(row=46, column=5, value="=SUM(E43:E45)")
mb.cell(row=46, column=6, value="=SUM(F43:F45)")
mb.cell(row=46, column=7, value="=SUM(G43:G45)")
mb.cell(row=46, column=8, value='=IF(E46>=$D$41*$D$4,"All income accounted for","Check inputs")')
box(mb, 43, 46, 2, 8, money=(4, 5, 7), pct=(3, 6), zebra=False)
for r in range(43, 47):
    mb.cell(row=r, column=2).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
    for c in (3, 6, 8):
        mb.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    for c in (4, 5, 7):
        mb.cell(row=r, column=c).alignment = Alignment(horizontal="right")
for c in range(2, 9):
    cell = mb.cell(row=46, column=c)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name=FONT, bold=True, color=WHITE, size=11)
status_cf(mb, "H43:H46", [
    ("On target", "D4EDDA", "1E7E34", True),
    ("Over target", "F8D7DA", "A02020", True),
    ("Behind plan", "FFF3CD", "997404", True),
    ("All income accounted for", "D1ECF1", "0C5460", False),
])
nb = mb.cell(row=48, column=2,
             value="Savings = income left after Needs and Wants (log any transfer with bucket "
                   "'Savings'). Target $ = income x months with actuals x target %.")
nb.font = Font(name=FONT, size=11, color=TEAL, italic=True)
mb.freeze_panes = "B8"

# ============================================================ SAVINGS GOALS
sg = sheet(wb, "SAVINGS GOALS")
g_headers = ["Goal", "Target", "Saved", "Remaining", "Progress %", "Target date",
             "Monthly needed", "Status"]
g_widths = [26, 14, 14, 14, 14, 15, 17, 15]
for i, (h, w) in enumerate(zip(g_headers, g_widths), start=1):
    sg.cell(row=1, column=i, value=h)
    sg.column_dimensions[get_column_letter(i)].width = w
hdr(sg, 1, g_headers)

goals = [
    ("Emergency fund", 12000, 8650, date(2026, 12, 31)),
    ("New laptop", 2400, 1780, date(2026, 10, 31)),
    ("Japan trip", 5000, 2150, date(2027, 4, 30)),
    ("Car down payment", 8000, 3300, date(2027, 6, 30)),
    ("UX design course", 1500, 1500, date(2026, 9, 30)),
    ("Home deposit", 40000, 11900, date(2029, 12, 31)),
]
for i, (name, target, saved, tdate) in enumerate(goals):
    r = 2 + i
    sg.cell(row=r, column=1, value=name)
    sg.cell(row=r, column=2, value=target)
    sg.cell(row=r, column=3, value=saved)
    sg.cell(row=r, column=4, value=f"=MAX(B{r}-C{r},0)")
    sg.cell(row=r, column=5, value=f"=IF(B{r}=0,0,C{r}/B{r})")
    sg.cell(row=r, column=6, value=tdate)
    sg.cell(row=r, column=7,
            value=f'=IF(D{r}<=0,0,ROUND(D{r}/MAX(1,(YEAR(F{r})-YEAR(TODAY()))*12'
                  f'+MONTH(F{r})-MONTH(TODAY())),0))')
    sg.cell(row=r, column=8,
            value=f'=IF(E{r}>=1,"Complete",IF(E{r}>=((MONTH(TODAY())-1)/12),'
                  f'"On track","Behind"))')
last_sg = 1 + len(goals)
box(sg, 2, last_sg, 1, 8, money=(2, 3, 4, 7), dates=(6,), pct=(5,))
for r in range(2, last_sg + 1):
    sg.cell(row=r, column=1).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
    sg.cell(row=r, column=5).alignment = Alignment(horizontal="center")
    sg.cell(row=r, column=8).alignment = Alignment(horizontal="center")
sg.freeze_panes = "A2"
sg.auto_filter.ref = f"A1:H{last_sg}"
sg.conditional_formatting.add(
    f"E2:E{last_sg}",
    ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                   mid_type="num", mid_value=0.5, mid_color="FFEB84",
                   end_type="num", end_value=1, end_color="63BE7B"))
sg.conditional_formatting.add(
    f"C2:C{last_sg}",
    DataBarRule(start_type="num", start_value=0, end_type="max", color=TEAL))
status_cf(sg, "H2:H100", [
    ("Complete", "D4EDDA", "1E7E34", True),
    ("On track", "D1ECF1", "0C5460", False),
    ("Behind", "F8D7DA", "A02020", True),
])

# ============================================================ DEBT SNOWBALL
ds = sheet(wb, "DEBT SNOWBALL")
d_headers = ["Creditor", "Balance", "APR", "Min payment", "Extra payment", "Payoff order",
             "Months to payoff", "Interest saved", "Status"]
d_widths = [22, 14, 10, 14, 14, 14, 17, 15, 18]
for i, (h, w) in enumerate(zip(d_headers, d_widths), start=1):
    ds.cell(row=1, column=i, value=h)
    ds.column_dimensions[get_column_letter(i)].width = w
hdr(ds, 1, d_headers)
ds.cell(row=1, column=11, value="Months (min only)")
ds.cell(row=1, column=12, value="Interest (min only)")
for c in (11, 12):
    cell = ds.cell(row=1, column=c)
    cell.font = Font(name=FONT, bold=True, color=WHITE, size=10)
    cell.fill = PatternFill("solid", fgColor=GREY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
    ds.column_dimensions[get_column_letter(c)].width = 18

debts = [
    ("Visa Classic", 3240, 0.229, 95, 0),
    ("Store card", 640, 0.275, 25, 250),
    ("Car loan", 9450, 0.064, 265, 0),
    ("Student loan", 14800, 0.052, 175, 0),
    ("Laptop instalments", 1180, 0.0, 99, 0),
    ("Personal loan", 2400, 0.119, 70, 0),
]
for i, (cred, bal, apr, mini, extra) in enumerate(debts):
    r = 2 + i
    ds.cell(row=r, column=1, value=cred)
    ds.cell(row=r, column=2, value=bal)
    ds.cell(row=r, column=3, value=apr)
    ds.cell(row=r, column=4, value=mini)
    ds.cell(row=r, column=5, value=extra)
    ds.cell(row=r, column=6,
            value=f'=IF(B{r}<=0,99,SUMPRODUCT(($B$2:$B$7<B{r})*1)+COUNTIF($B$2:B{r},B{r}))')
    pay = f"(D{r}+E{r})"
    ds.cell(row=r, column=7,
            value=f'=IF({pay}<=0,999,IF(ROUND(B{r}*(C{r}/12),2)>={pay},999,'
                  f'IF(C{r}=0,ROUNDUP(B{r}/{pay},0),'
                  f'ROUND(-LN(1-B{r}*(C{r}/12)/{pay})/LN(1+C{r}/12),0))))')
    ds.cell(row=r, column=8,
            value=f'=IF(OR(K{r}>=999,G{r}>=999),0,MAX(0,ROUND(D{r}*K{r}-{pay}*G{r},0)))')
    ds.cell(row=r, column=9,
            value=f'=IF(B{r}<=0,"Paid off",IF(F{r}=1,"Attack first",'
                  f'IF(G{r}>=999,"Never shrinks","Snowball later")))')
    ds.cell(row=r, column=11,
            value=f'=IF(D{r}<=0,999,IF(ROUND(B{r}*(C{r}/12),2)>=D{r},999,'
                  f'IF(C{r}=0,ROUNDUP(B{r}/D{r},0),'
                  f'ROUND(-LN(1-B{r}*(C{r}/12)/D{r})/LN(1+C{r}/12),0))))')
    ds.cell(row=r, column=12,
            value=f'=IF(K{r}>=999,0,MAX(0,ROUND(D{r}*K{r}-B{r},0)))')
last_ds = 1 + len(debts)
box(ds, 2, last_ds, 1, 9, money=(2, 4, 5, 8), zebra=True)
box(ds, 2, last_ds, 11, 12, money=(12,), zebra=False)
for r in range(2, last_ds + 1):
    ds.cell(row=r, column=1).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
    ds.cell(row=r, column=3).number_format = "0.0%"
    for c in (3, 6, 7):
        ds.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    ds.cell(row=r, column=9).alignment = Alignment(horizontal="center")
    for c in (11, 12):
        ds.cell(row=r, column=c).font = Font(name=FONT, size=9, color=GREY)
        ds.cell(row=r, column=c).alignment = Alignment(horizontal="center")
ds.freeze_panes = "A2"
ds.auto_filter.ref = f"A1:I{last_ds}"
dv_apr = DataValidation(type="decimal", operator="between", formula1="0", formula2="0.6",
                        allow_blank=True)
dv_apr.error = "APR must be between 0% and 60%"
ds.add_data_validation(dv_apr)
dv_apr.add("C2:C100")
status_cf(ds, "I2:I100", [
    ("Attack first", "D4EDDA", "1E7E34", True),
    ("Paid off", "D1ECF1", "0C5460", True),
    ("Never shrinks", "F8D7DA", "A02020", True),
    ("Snowball later", "FFF3CD", "997404", False),
])
ds.conditional_formatting.add(
    f"B2:B{last_ds}",
    DataBarRule(start_type="num", start_value=0, end_type="max", color=GOLD))
ds["A9"] = "Smallest balance first - put every spare dollar (extra payment) on the top row."
ds["A9"].font = Font(name=FONT, size=11, color=TEAL, italic=True)
ds["A10"] = "Columns K and L are helper workings for the interest comparison."
ds["A10"].font = Font(name=FONT, size=10, color=GREY, italic=True)

# ================================================================ NET WORTH
nw = sheet(wb, "NET WORTH")
n_headers = ["Item", "Type", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
             "Latest (Jun)", "Change vs Jan"]
n_widths = [22, 12, 12, 12, 12, 12, 12, 12, 15, 16]
for i, (h, w) in enumerate(zip(n_headers, n_widths), start=1):
    nw.cell(row=1, column=i, value=h)
    nw.column_dimensions[get_column_letter(i)].width = w
hdr(nw, 1, n_headers)

assets = [
    ("Cash (wallet)", [340, 310, 295, 350, 320, 305]),
    ("Checking", [2150, 2380, 2020, 2610, 2290, 2480]),
    ("Savings account", [6200, 6850, 7400, 8050, 8600, 9150]),
    ("Investments", [11400, 11950, 12300, 12850, 13200, 13650]),
    ("Crypto", [2750, 3120, 2980, 3450, 3680, 3900]),
]
liabs = [
    ("Credit cards", [3850, 3620, 3400, 3050, 2680, 2300]),
    ("Car loan", [9450, 9185, 8920, 8655, 8390, 8125]),
    ("Student loan", [14800, 14680, 14560, 14440, 14320, 14200]),
]
band(nw, 2, 1, 10, "ASSETS", fill=TEAL)
r = 3
asset_rows = []
for name, vals in assets:
    nw.cell(row=r, column=1, value=name)
    nw.cell(row=r, column=2, value="Asset")
    for i, v in enumerate(vals):
        nw.cell(row=r, column=3 + i, value=v)
    asset_rows.append(r)
    r += 1
band(nw, r, 1, 10, "LIABILITIES", fill=GOLD, color=NAVY)
r += 1
liab_rows = []
for name, vals in liabs:
    nw.cell(row=r, column=1, value=name)
    nw.cell(row=r, column=2, value="Liability")
    for i, v in enumerate(vals):
        nw.cell(row=r, column=3 + i, value=v)
    liab_rows.append(r)
    r += 1
nw_row = r
nw.cell(row=nw_row, column=1, value="NET WORTH")
nw.cell(row=nw_row, column=2, value="Net")
first_a, last_a = asset_rows[0], asset_rows[-1]
first_l, last_l = liab_rows[0], liab_rows[-1]
for c in range(3, 9):
    letter = get_column_letter(c)
    nw.cell(row=nw_row, column=c,
            value=f"=SUM({letter}{first_a}:{letter}{last_a})-SUM({letter}{first_l}:{letter}{last_l})")
all_rows = list(asset_rows) + list(liab_rows)
for rr in all_rows + [nw_row]:
    nw.cell(row=rr, column=9, value=f"=H{rr}")
    nw.cell(row=rr, column=10, value=f"=H{rr}-C{rr}")
box(nw, first_a, nw_row, 1, 10, money=tuple(range(3, 11)), zebra=True)
for rr in all_rows:
    nw.cell(row=rr, column=1).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
    nw.cell(row=rr, column=2).alignment = Alignment(horizontal="center")
    nw.cell(row=rr, column=2).font = Font(name=FONT, size=10, color=GREY)
for c in range(1, 11):
    cell = nw.cell(row=nw_row, column=c)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name=FONT, bold=True, size=12, color=WHITE)
    cell.alignment = Alignment(horizontal="center" if c > 2 else "left")
    if c > 2:
        cell.number_format = MONEY
nw.row_dimensions[nw_row].height = 24

grow_row = nw_row + 2
nw.cell(row=grow_row, column=1, value="Net worth growth since January")
nw.cell(row=grow_row, column=1).font = Font(name=FONT, bold=True, size=11, color=NAVY)
nw.cell(row=grow_row, column=9, value=f"=J{nw_row}")
nw.cell(row=grow_row, column=10, value=f"=IF(C{nw_row}=0,0,J{nw_row}/ABS(C{nw_row}))")
for c in (9, 10):
    cell = nw.cell(row=grow_row, column=c)
    cell.border = border
    cell.fill = PatternFill("solid", fgColor=LIGHT)
    cell.font = Font(name=FONT, bold=True, size=12, color=TEAL)
    cell.alignment = Alignment(horizontal="right")
nw.cell(row=grow_row, column=9).number_format = MONEY
nw.cell(row=grow_row, column=10).number_format = PCT
nw.conditional_formatting.add(
    f"J{first_a}:J{nw_row}",
    CellIsRule(operator="lessThan", formula=["0"],
               font=Font(name=FONT, bold=True, color="A02020")))
nw.conditional_formatting.add(
    f"J{first_a}:J{nw_row}",
    CellIsRule(operator="greaterThan", formula=["0"],
               font=Font(name=FONT, bold=True, color="1E7E34")))
note_nw = nw.cell(row=grow_row + 2, column=1,
                  value="Add Jul to Dec in the same pattern - copy the last month column across.")
note_nw.font = Font(name=FONT, size=11, color=TEAL, italic=True)
nw.freeze_panes = "C2"

# ================================================================ DASHBOARD
db = sheet(wb, "DASHBOARD")
widths(db, {"A": 3, "B": 30, "C": 17, "D": 30, "E": 17, "F": 3,
            "G": 10, "H": 16, "I": 28})
db["B2"] = "BUDGET DASHBOARD"
db["B2"].font = Font(name=FONT, bold=True, size=18, color=NAVY)
db["B3"] = "One page, live formulas. Change the reporting month in cell C4."
db["B3"].font = Font(name=FONT, size=11, color=TEAL, italic=True)

lb3 = db.cell(row=4, column=2, value="Reporting month:")
lb3.font = Font(name=FONT, bold=True, size=11, color=GREY)
lb3.border = border
rm = db.cell(row=4, column=3, value=date(2026, 5, 1))
rm.font = Font(name=FONT, bold=True, size=12, color=TEAL)
rm.fill = PatternFill("solid", fgColor=LIGHT)
rm.border = border
rm.number_format = DATEF
rm.alignment = Alignment(horizontal="center")

MB = "'MONTHLY BUDGET'"
EL2 = "'EXPENSE LOG'"
SG = "'SAVINGS GOALS'"
DS = "'DEBT SNOWBALL'"
NW = "'NET WORTH'"
spent_month = (f'=SUMIFS({EL2}!$F:$F,{EL2}!$A:$A,">="&$C$4,'
               f'{EL2}!$A:$A,"<"&DATE(YEAR($C$4),MONTH($C$4)+1,1))')

kpi(db, 6, 2, "INCOME (reporting month)", f"={MB}!$D$41")
kpi(db, 9, 2, "SPENT (reporting month)", spent_month)
kpi(db, 12, 2, "SAVED (income - spent)", "=B7-B10")
kpi(db, 15, 2, "SAVINGS RATE", "=IF(B7=0,0,B13/B7)", fmt=PCT)
kpi(db, 18, 2, "AVG GOALS PROGRESS", f"=SUMPRODUCT({SG}!$E$2:$E$7)/6", fmt=PCT)
kpi(db, 21, 2, "DEBT-FREE DATE (last debt)",
    f"=DATE(YEAR(TODAY()),MONTH(TODAY())+SUMPRODUCT(MAX({DS}!$G$2:$G$7)),DAY(TODAY()))",
    fmt=DATEF)

kpi(db, 6, 4, "BIGGEST CATEGORY (month)",
    '=IF(SUM($H$5:$H$18)=0,"-",VLOOKUP(SUMPRODUCT(MAX($H$5:$H$18)),$G$5:$I$18,3,FALSE))',
    fmt="General")
kpi(db, 9, 4, "CATEGORIES OVER PLAN (#)",
    f"=SUMPRODUCT(({MB}!$U$25:$U$38>1)*1)", fmt="#,##0")
kpi(db, 12, 4, "NET WORTH (latest)", f"={NW}!$I${nw_row}")
kpi(db, 15, 4, "NET WORTH CHANGE SINCE JAN", f"={NW}!$J${nw_row}")
kpi(db, 18, 4, "BUDGET HEALTH",
    f'=IF(SUM({MB}!$U$25:$U$38)=0,"No data yet",'
    f'IF(SUMPRODUCT(({MB}!$U$25:$U$38>1)*1)=0,"Within plan","Over plan"))', fmt="General")
kpi(db, 21, 4, "DEBT SNOWBALL EXTRAS ($/mo)",
    f"=SUM({DS}!$E$2:$E$7)")

# Top 5 spending categories for the reporting month
band(db, 25, 2, 4, "TOP 5 SPENDING CATEGORIES (reporting month)")
hdr(db, 26, ["Category", "Spent", "Share of month"], start=2, fill=TEAL)
for i in range(5):
    r = 27 + i
    db.cell(row=r, column=2, value=f"=VLOOKUP({i + 1},$G$5:$I$18,3,FALSE)")
    db.cell(row=r, column=3, value=f"=VLOOKUP({i + 1},$G$5:$I$18,2,FALSE)")
    db.cell(row=r, column=4, value=f'=IF($B$10=0,0,C{r}/$B$10)')
box(db, 27, 31, 2, 4, money=(3,))
for r in range(27, 32):
    db.cell(row=r, column=2).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
    db.cell(row=r, column=3).alignment = Alignment(horizontal="right")
    db.cell(row=r, column=4).number_format = PCT
    db.cell(row=r, column=4).alignment = Alignment(horizontal="center")

# Income vs spend by month
band(db, 33, 2, 5, "INCOME VS SPEND BY MONTH - 2026")
hdr(db, 34, ["Month", "Income", "Spent", "Saved"], start=2, fill=TEAL)
for m in range(1, 13):
    r = 34 + m
    db.cell(row=r, column=2, value=f"{MONTHS_SHORT[m - 1]} 2026")
    db.cell(row=r, column=3, value=f"={MB}!$D$41")
    db.cell(row=r, column=4,
            value=f'=SUMIFS({EL2}!$F:$F,{EL2}!$A:$A,">="&DATE(2026,{m},1),'
                  f'{EL2}!$A:$A,"<"&DATE(2026,{m + 1},1))')
    db.cell(row=r, column=5, value=f"=C{r}-D{r}")
box(db, 35, 46, 2, 5, money=(3, 4, 5))
for r in range(35, 47):
    db.cell(row=r, column=2).font = Font(name=FONT, bold=True, size=11, color="2B2D42")
db.cell(row=47, column=2, value="YEAR TOTAL")
for c in range(3, 6):
    letter = get_column_letter(c)
    db.cell(row=47, column=c, value=f"=SUM({letter}35:{letter}46)")
for c in range(2, 6):
    cell = db.cell(row=47, column=c)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name=FONT, bold=True, color=WHITE, size=11)
    cell.border = border
    cell.alignment = Alignment(horizontal="right" if c > 2 else "left")
    if c > 2:
        cell.number_format = MONEY
db.conditional_formatting.add(
    "E35:E46",
    CellIsRule(operator="lessThan", formula=["0"],
               font=Font(name=FONT, bold=True, color="A02020")))

db["B49"] = "All numbers come from EXPENSE LOG, MONTHLY BUDGET, SAVINGS GOALS, DEBT SNOWBALL " \
            "and NET WORTH - nothing is typed here."
db["B49"].font = Font(name=FONT, size=11, color=TEAL, italic=True)

# Category ranking workings (feeds Biggest Category + Top 5)
db["G3"] = "CATEGORY RANKINGS (auto)"
db["G3"].font = Font(name=FONT, bold=True, size=10, color=GREY)
hdr(db, 4, ["Rank", "Spent", "Category"], start=7, fill=GREY)
for i, cat in enumerate(CATEGORIES):
    r = 5 + i
    db.cell(row=r, column=7,
            value=f"=SUMPRODUCT(($H$5:$H$18>H{r})*1)+COUNTIF($H$5:H{r},H{r})")
    db.cell(row=r, column=8,
            value=f'=SUMIFS({EL2}!$F:$F,{EL2}!$B:$B,$I{r},{EL2}!$A:$A,">="&$C$4,'
                  f'{EL2}!$A:$A,"<"&DATE(YEAR($C$4),MONTH($C$4)+1,1))')
    db.cell(row=r, column=9, value=cat)
box(db, 5, 18, 7, 9, money=(8,), zebra=False)
for r in range(5, 19):
    for c in (7, 8, 9):
        db.cell(row=r, column=c).font = Font(name=FONT, size=9, color=GREY)
    db.cell(row=r, column=7).alignment = Alignment(horizontal="center")
    db.cell(row=r, column=8).alignment = Alignment(horizontal="right")

# ---------------------------------------------------------------- SAVE
wb._sheets = [wb["Start Here"], wb["DASHBOARD"], wb["MONTHLY BUDGET"], wb["EXPENSE LOG"],
              wb["SAVINGS GOALS"], wb["DEBT SNOWBALL"], wb["NET WORTH"]]
wb.active = 0
os.makedirs(os.path.dirname(os.path.abspath(__file__)), exist_ok=True)
out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Personal_Budget_and_Savings_Tracker_2026.xlsx")
wb.save(out)
print("saved:", out)
